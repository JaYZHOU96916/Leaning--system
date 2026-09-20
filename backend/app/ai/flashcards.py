import asyncio
import json
import re
from collections.abc import Awaitable, Callable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
from pydantic import BaseModel, Field
from sqlmodel import Session, select

from app.core.config import Settings
from app.models import Flashcard, FlashcardJob, Material


class FlashcardDraft(BaseModel):
    question: str = Field(min_length=1)
    answer: str = Field(min_length=1)
    source: str | None = None


LLMCall = Callable[[str], Awaitable[str]]


def _strip_code_fence(value: str) -> str:
    return re.sub(r"^```(?:json)?\s*|\s*```$", "", value.strip(), flags=re.IGNORECASE)


def parse_flashcard_payload(raw: str, *, source: str | None = None) -> list[FlashcardDraft]:
    payload: Any = json.loads(_strip_code_fence(raw))
    if isinstance(payload, dict):
        payload = payload.get("cards", [])
    if not isinstance(payload, list):
        raise ValueError("LLM response must be a JSON array or an object containing cards")
    cards: list[FlashcardDraft] = []
    seen: set[tuple[str, str]] = set()
    for item in payload:
        if not isinstance(item, dict):
            continue
        card = FlashcardDraft(
            question=str(item.get("question", "")).strip(),
            answer=str(item.get("answer", "")).strip(),
            source=item.get("source") or source,
        )
        key = (card.question.casefold(), card.answer.casefold())
        if key not in seen:
            seen.add(key)
            cards.append(card)
    return cards


class LLMFlashcardExtractor:
    def __init__(self, llm_call: LLMCall, *, max_cards: int = 30):
        self.llm_call = llm_call
        self.max_cards = max_cards

    async def extract(self, text: str, *, source: str | None = None) -> list[FlashcardDraft]:
        prompt = (
            "Extract the most important academic concepts from the following material. "
            "Return strict JSON with a `cards` array; each item must have `question`, "
            "`answer`, and optional `source`. Keep answers concise and do not invent facts.\n\n"
            f"MATERIAL:\n{text[:50000]}"
        )
        return parse_flashcard_payload(await self.llm_call(prompt), source=source)[: self.max_cards]


class HeuristicFlashcardExtractor:
    def __init__(self, *, max_cards: int = 10):
        self.max_cards = max_cards

    async def extract(self, text: str, *, source: str | None = None) -> list[FlashcardDraft]:
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
        cards = [
            FlashcardDraft(
                question=f"What is the key idea in section {index}?",
                answer=paragraph,
                source=source,
            )
            for index, paragraph in enumerate(paragraphs, start=1)
        ]
        return cards[: self.max_cards]


class OpenAICompatibleLLM:
    def __init__(self, settings: Settings, *, client: httpx.AsyncClient | None = None):
        if not settings.llm_api_key or not settings.llm_api_key.get_secret_value():
            raise ValueError("LLM_API_KEY is required for OpenAI-compatible extraction")
        self.settings = settings
        self.client = client

    async def __call__(self, prompt: str) -> str:
        payload = {
            "model": self.settings.llm_model,
            "temperature": 0,
            "messages": [{"role": "user", "content": prompt}],
            "response_format": {"type": "json_object"},
        }
        headers = {"Authorization": f"Bearer {self.settings.llm_api_key.get_secret_value()}"}
        url = f"{self.settings.llm_base_url.rstrip('/')}/chat/completions"
        if self.client is not None:
            response = await self.client.post(url, json=payload, headers=headers)
        else:
            async with httpx.AsyncClient(timeout=self.settings.llm_timeout_seconds) as client:
                response = await client.post(url, json=payload, headers=headers)
        response.raise_for_status()
        return str(response.json()["choices"][0]["message"]["content"])


def extractor_for_settings(settings: Settings):
    if settings.llm_api_key and settings.llm_api_key.get_secret_value():
        return LLMFlashcardExtractor(OpenAICompatibleLLM(settings))
    return HeuristicFlashcardExtractor()


def extract_text_from_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".html", ".htm"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        from docx import Document

        return "\n\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    if suffix == ".pptx":
        from pptx import Presentation

        parts: list[str] = []
        for slide in Presentation(path).slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    parts.append(shape.text)
        return "\n\n".join(parts)
    raise ValueError(f"Unsupported material type: {suffix or 'unknown'}")


def to_anki_tsv(cards: list[FlashcardDraft]) -> str:
    rows = []
    for card in cards:
        question = card.question.replace("\t", " ").replace("\n", "<br>")
        answer = card.answer.replace("\t", " ").replace("\n", "<br>")
        rows.append(f"{question}\t{answer}")
    return "\n".join(rows) + ("\n" if rows else "")


async def generate_flashcards_for_job(
    job_id: int,
    *,
    settings: Settings,
    session_factory: Callable[[], Session],
) -> int:
    with session_factory() as session:
        job = session.get(FlashcardJob, job_id)
        if job is None:
            raise ValueError("Flashcard job not found")
        material = session.get(Material, job.material_id)
        if material is None or not material.local_path:
            job.status = "failed"
            job.error = "Material file is unavailable"
            session.commit()
            return 0
        job.status = "running"
        session.commit()
        path = Path(material.local_path)
        course_id = job.course_id
        material_id = material.id
        material_name = material.name
    try:
        text = await asyncio.to_thread(extract_text_from_file, path)
        cards = await extractor_for_settings(settings).extract(text, source=material_name)
    except Exception as exc:
        with session_factory() as session:
            job = session.get(FlashcardJob, job_id)
            if job:
                job.status = "failed"
                job.error = str(exc)[:500]
                session.commit()
        return 0
    with session_factory() as session:
        created = 0
        for card in cards:
            exists = session.exec(
                select(Flashcard)
                .where(Flashcard.material_id == material_id)
                .where(Flashcard.question == card.question)
            ).first()
            if exists is None:
                session.add(
                    Flashcard(
                        course_id=course_id,
                        material_id=material_id,
                        question=card.question,
                        answer=card.answer,
                        source=card.source,
                    )
                )
                created += 1
        job = session.get(FlashcardJob, job_id)
        if job:
            job.status = "completed"
            job.cards_created = created
            job.completed_at = datetime.now(UTC)
        session.commit()
    return created
